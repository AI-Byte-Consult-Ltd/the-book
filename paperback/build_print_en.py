#!/usr/bin/env python3
"""Build the two-volume English paperback manuscript (6x9 in) of «The Appearance» for Amazon KDP.
Volume I = Chapter I ... Chapter XXII; Volume II = Chapter XXIII ... Reality XI + About the Author.
Output: paperback/The Appearance. Volume I|II .docx and .pdf"""
import re, subprocess, pathlib, shutil, tempfile, copy
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

here = pathlib.Path(__file__).resolve().parent
src = (here.parent / "epub" / "book_en.md").read_text(encoding="utf-8")
SPLIT_AFTER = "Chapter XXII. The Rose Garden"      # end of Volume I (~50% of the text)

front_all = src[:src.index("# Part One.")]
about_book = front_all[front_all.index("# About the Book"):]
copyright_txt = front_all[front_all.index("# Copyright"):front_all.index("# About the Book")]
copyright_txt = copyright_txt.replace("# Copyright {#copyright}", "# Copyright")
copyright_txt = re.sub(r"\n\nISBN-13:[^\n]*", "", copyright_txt)       # each volume needs its own ISBN
about_author = src[src.index("# About the Author"):]
body = src[src.index("# Part One."):src.index("# About the Author")]
chunks = re.split(r"^(?=## )", body, flags=re.M)[1:]
heads = [re.match(r"## (.+)", c).group(1) for c in chunks]
cut = heads.index(SPLIT_AFTER) + 1
vols = {1: chunks[:cut], 2: chunks[cut:]}

def scene(t): return re.sub(r"^---\s*$", '\n::: {custom-style="Scene Break"}\n\\\\* \\\\* \\\\*\n:::\n', t, flags=re.M)

def volume_md(n, toc_pages):
    items = vols[n]; hs = [re.match(r"## (.+)", c).group(1) for c in items]
    toc = "# Contents\n\n" + "\n\n".join(f'::: {{custom-style="TOC Entry"}}\n{h}\t{toc_pages.get(h, 0)}\n:::' for h in hs) + "\n\n"
    roman = "I" * n if n < 3 else "II"
    title = ('::: {custom-style="Book Title"}\nThe Appearance\n:::\n\n'
             f'::: {{custom-style="Book Volume"}}\nVolume {"I" if n == 1 else "II"}\n:::\n\n'
             '::: {custom-style="Book Subtitle"}\nA Mystical Adventure Novel\n:::\n\n'
             '::: {custom-style="Book Author"}\nAlexander Lunin\n:::\n\n'
             '::: {custom-style="Book Subtitle"}\nSeries "Terra Incognita". Book One\n:::\n\n')
    front = title + scene(copyright_txt) + "\n" + (scene(about_book) + "\n" if n == 1 else "") + toc
    text = "".join(items)
    text = re.sub(r"^## (.+)$", r"## \1", text, flags=re.M)
    tail = ""
    if n == 1:
        tail = '\n\n::: {custom-style="Scene Break"}\nEnd of Volume I. The story continues in Volume II.\n:::\n'
    else:
        tail = "\n\n" + about_author
    return front + "\n<!--BODY-->\n" + scene(text) + tail, hs

def set_footer_page_number(section):
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for t, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        r = p.add_run(); r.font.size = Pt(10); r.font.name = "Times New Roman"
        if t:
            f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), t); r._r.append(f)
        else:
            i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = txt; r._r.append(i)

def build(n, toc_pages):
    md, hs = volume_md(n, toc_pages)
    marker = "<!--BODY-->"
    first_head_text = hs[0]
    tmp = pathlib.Path(tempfile.mkdtemp()) / "v.md"
    tmp.write_text(md.replace(marker, ""), encoding="utf-8")
    out = here / f"The Appearance. Volume {'I' if n == 1 else 'II'}.docx"
    subprocess.run(["pandoc", str(tmp), "-f", "markdown+smart", "-t", "docx", "-M", "lang=en-US", "-o", str(out)], check=True)
    d = Document(out)
    sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(6), Inches(9)
    sec.left_margin, sec.right_margin = Inches(0.85), Inches(0.65)     # inside / outside (mirrored)
    sec.top_margin, sec.bottom_margin = Inches(0.75), Inches(0.8)
    sec.footer_distance = Inches(0.4)
    mm = OxmlElement("w:mirrorMargins"); d.settings.element.insert(0, mm)
    def style(name, **kw):
        try: s = d.styles[name]
        except KeyError: s = d.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        s.font.name = "Times New Roman"; s.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")
        s.font.color.rgb = RGBColor(0, 0, 0); s.font.size = Pt(kw.get("size", 11)); s.font.bold = kw.get("bold", False); s.font.italic = False
        f = s.paragraph_format; f.alignment = kw.get("align", WD_ALIGN_PARAGRAPH.JUSTIFY)
        f.first_line_indent = kw.get("indent"); f.left_indent = None
        f.space_before = Pt(kw.get("before", 0)); f.space_after = Pt(kw.get("after", 0)); f.line_spacing = kw.get("ls", 1.15)
        f.page_break_before = kw.get("pb", False); f.keep_with_next = kw.get("kwn", False)
        f.widow_control = True
        return s
    C, L = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT
    for nm in ("Normal", "Body Text", "First Paragraph", "Compact"): style(nm, indent=Inches(0.25))
    style("Heading 1", size=18, bold=True, align=C, before=60, after=20, pb=True, kwn=True, indent=None)
    style("Heading 2", size=15, bold=True, align=C, before=90, after=24, pb=True, kwn=True, indent=None)
    style("Book Title", size=30, bold=True, align=C, before=130, after=10, indent=None)
    style("Book Volume", size=18, align=C, before=6, after=24, indent=None)
    style("Book Subtitle", size=13, align=C, before=6, after=6, indent=None)
    style("Book Author", size=17, align=C, before=30, after=30, indent=None)
    style("Scene Break", align=C, before=12, after=12, indent=None)
    toc_style = style("TOC Entry", align=L, after=2, indent=None, ls=1.0, size=10.5)
    toc_style.paragraph_format.tab_stops.add_tab_stop(Inches(4.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    # unindented front matter (copyright)
    inside = False
    for p in d.paragraphs:
        if p.style.name == "Heading 1": inside = p.text == "Copyright"
        elif inside: p.paragraph_format.first_line_indent = Inches(0); p.alignment = L; p.paragraph_format.space_after = Pt(6)
    # section break before the first body heading: front matter has no page numbers, body restarts at 1
    first = next(p for p in d.paragraphs if p.style.name == "Heading 2" and p.text == first_head_text)
    prev = first._p.getprevious()
    while prev.tag != qn("w:p"): prev = prev.getprevious()
    body_sectPr = d.element.body.sectPr
    front_sectPr = copy.deepcopy(body_sectPr)
    for tag in ("w:footerReference", "w:headerReference"):
        for e in front_sectPr.findall(qn(tag)): front_sectPr.remove(e)
    pPr = prev.get_or_add_pPr() if hasattr(prev, "get_or_add_pPr") else prev.find(qn("w:pPr"))
    if pPr is None:
        pPr = OxmlElement("w:pPr"); prev.insert(0, pPr)
    pPr.append(front_sectPr)
    sections = d.sections
    body_sec = sections[-1]
    set_footer_page_number(body_sec)
    pg = OxmlElement("w:pgNumType"); pg.set(qn("w:start"), "1"); body_sec._sectPr.append(pg)
    d.core_properties.title = f"The Appearance. Volume {'I' if n == 1 else 'II'}"; d.core_properties.author = "Alexander Lunin"
    d.save(out)
    return out, hs

def to_pdf(docx):
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(here), str(docx)], check=True,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    return docx.with_suffix(".pdf")

def page_map(pdf, hs):
    n = int(re.search(r"Pages:\s+(\d+)", subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True).stdout).group(1))
    pages = {}
    for pg in range(1, n + 1):
        t = subprocess.run(["pdftotext", "-f", str(pg), "-l", str(pg), "-layout", str(pdf), "-"], capture_output=True, text=True).stdout
        lines = [l.strip().replace("\u2019", "'").replace("\u201c", '"').replace("\u201d", '"') for l in t.splitlines() if l.strip()]
        cands = [lines[0]] + ([lines[0] + " " + lines[1], lines[0] + lines[1]] if len(lines) > 1 else []) if lines else []
        for c in cands:
            if c in hs and c not in pages: pages[c] = pg
    first = pages[hs[0]]
    return {h: pages[h] - first + 1 for h in hs}, n

for v in (1, 2):
    docx, hs = build(v, {})
    pdf = to_pdf(docx)
    pm, _ = page_map(pdf, hs)
    docx, hs = build(v, pm)            # second pass with real page numbers in Contents
    pdf = to_pdf(docx)
    pm2, total = page_map(pdf, hs)
    assert pm == pm2, "page numbers shifted after filling Contents"
    print(docx.name, "- PDF pages:", total)
